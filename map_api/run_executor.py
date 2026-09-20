"""Satellite DAG handlers; AgentRun checkpoints are the execution authority."""
import base64
import json
import uuid
from collections import defaultdict
from datetime import date, timedelta
from io import BytesIO

import numpy as np
import requests
from django.db import transaction
from django.utils import timezone
from jsonschema import Draft202012Validator

from .agent.providers import ProviderError, configured_provider
from .data_contract import SOURCE_CAPABILITIES, build_data_requirements
from .imagery_sources import get_provider_for_collection
from .imagery_sources.earth_search import get_collection_profile
from .models import AgentRun, AgentSession, AnalysisRun, Evidence, Finding
from .remote_sensing_indices import INDEX_DEFINITIONS
from .run_journal import append_locked, load_checkpoint, lock_run
from .run_kernel import _set_state
from .run_products import (apply_masks, calculate_products, compare_products, geotiff_bytes, load_arrays,
                           read_product, read_scene_assets, rgb_preview, save_arrays, write_product)
from .run_scheduler import (StepConflict, _owned_step, claim_step, fail_step, finish_step, renew_step, skip_step)


PLAN_SPEC = [
    ("understand_goal", "理解调查目标", [], "slots"),
    ("resolve_aoi", "定位调查范围", ["understand_goal"], "aoi"),
    ("declare_data_requirements", "声明数据需求", ["resolve_aoi"], "requirements"),
    ("match_source_capabilities", "匹配数据源能力", ["declare_data_requirements"], "capabilities"),
    ("search_scenes", "检索影像场景", ["match_source_capabilities"], "candidates"),
    ("quality_gate", "校验日期、云量和覆盖", ["search_scenes"], "quality"),
    ("select_product", "选择合格产品", ["quality_gate"], "periods"),
    ("read_assets", "读取并校准原始波段", ["select_product"], "assets"),
    ("apply_aoi_and_qa_mask", "应用 AOI 与 QA 掩膜", ["read_assets"], "mask_statistics"),
    ("compute_metric", "计算遥感指标", ["apply_aoi_and_qa_mask"], "metrics"),
    ("compare_temporal_scenes", "比较两期共同有效像元", ["compute_metric"], "changes"),
    ("visual_review", "视觉复核", ["compute_metric"], "visual_review"),
    ("evidence_compilation", "整理证据与限制", ["visual_review", "compare_temporal_scenes"], "evidence_refs"),
    ("final_review", "复核最终结论", ["evidence_compilation"], "final"),
    ("publish_artifacts", "生成报告与结果产物", ["final_review"], "report"),
]


def default_plan():
    return [{"step_id": sid, "kind": sid, "label": label, "purpose": label,
             "depends_on": dependencies, "max_attempts": 3,
             "optional": sid in {"compare_temporal_scenes"},
             "retry_policy": {"recoverable": True, "automatic": False},
             "failure_conditions": ["必需输入缺失", "输出未通过 Schema 或质量门禁", "外部服务不可用"],
             "completion_schema": {"type": "object", "required": [field], "properties": {field: {}}}}
            for sid, label, dependencies, field in PLAN_SPEC]


SLOT_SCHEMA = {
    "type": "object", "additionalProperties": False,
    "required": ["task", "source", "place_name", "indices", "two_dates", "physical_measurement", "area_ratio"],
    "properties": {
        "task": {"enum": ["water", "vegetation", "agriculture", "land_use", "built_up", "small_target"]},
        "source": {"enum": ["sentinel2", "mapbox", "tianditu", "esri", "sentinel1", "copdem", "landsat"]}, "place_name": {"type": ["string", "null"]},
        "indices": {"type": "array", "uniqueItems": True, "items": {"enum": ["ndwi", "ndvi", "mndwi"]}},
        "two_dates": {"type": "boolean"}, "physical_measurement": {"type": "boolean"},
        "area_ratio": {"type": "boolean"}, "recent_acquisition": {"type": "boolean"},
        **{key: {"type": ["string", "null"], "pattern": "^\\d{4}-\\d{2}-\\d{2}$"}
           for key in ("date_start", "date_end", "before_date_start", "before_date_end")},
    },
}


class DomainBlocked(ValueError):
    def __init__(self, message, status="blocked"):
        self.status = status
        super().__init__(message)


# DAG 槽位 source → Earth Search collection 映射;sentinel1/copdem 走同一检索管线。
SOURCE_TO_COLLECTION = {
    "sentinel2": "sentinel-2-l2a",
    "sentinel1": "sentinel-1-grd",
    "copdem": "cop-dem-glo-30",
    "landsat": "landsat-c2-l2",
}


def _artifact(claim, product, title, kind, mime, *, refs=None):
    artifact_id = f"file:{claim.lease_plan_version}:{product['file_name']}"
    return {"artifact_id": artifact_id, "kind": kind, "title": title,
            "uri": f"/api/v2/agent/runs/{claim.run_id}/artifacts/{artifact_id}/download/",
            "mime_type": mime, "metadata": product, "evidence_refs": refs or []}


def _evidence(claim, *, kind, scene_id, value, metric="", method="", contract=None, mask=None, aoi=None, asset_id="", limitations=None):
    return {"evidence_id": f"v{claim.lease_plan_version}:{claim.step_id}:{scene_id}:{metric or kind}",
            "kind": kind, "scene_id": str(scene_id), "metric": metric, "value": value,
            "asset_id": asset_id, "method": method, "data_contract": contract or {},
            "mask_statistics": mask or {}, "aoi": aoi,
            "limitations": limitations or ["公开遥感数据仅作区域筛查"]}


class SatelliteHandlers:
    def __init__(self, run, claim, context, provider_factory=configured_provider):
        self.run, self.claim, self.context = run, claim, context
        self.files = []
        self.evidence = []
        self.artifacts = []
        self.provider_factory = provider_factory

    def before_request(self):
        renew_step(self.claim, lease_seconds=600)

    def provider(self):
        return self.provider_factory(name=self.run.provider, model=self.run.model,
                                     before_request=self.before_request, telemetry=self.telemetry)

    def telemetry(self, payload):
        with transaction.atomic():
            run = lock_run(self.run.id)
            _owned_step(run, self.claim)
            append_locked(run, "provider.requested", payload, command_key=uuid.uuid4().hex)

    def save_file(self, product, title, kind, mime, refs=None):
        self.files.append(product)
        artifact = _artifact(self.claim, product, title, kind, mime, refs=refs)
        self.artifacts.append(artifact)
        return artifact["artifact_id"]

    def understand_goal(self):
        slots = self.provider().plan({"goal": self.run.goal, "user_conditions": self.context.get("input", {}),
                                      "today": date.today().isoformat(),
                                      "instruction": "识别目标要求，不降低原始目标。水体必须 NDWI，植被必须 NDVI；明确两期变化需两个日期区间。"}, SLOT_SCHEMA)
        conditions = self.context.get("input", {})
        slots.update({key: value for key, value in conditions.items() if key in SLOT_SCHEMA["properties"]})
        Draft202012Validator(SLOT_SCHEMA).validate(slots)
        if slots["task"] == "water" and "ndwi" not in slots["indices"]:
            slots["indices"].append("ndwi")
        if slots["task"] in {"vegetation", "agriculture"} and "ndvi" not in slots["indices"]:
            slots["indices"].append("ndvi")
        if slots["source"] in ("sentinel2", "sentinel1", "landsat"):
            slots["date_end"] = slots.get("date_end") or date.today().isoformat()
            slots["date_start"] = slots.get("date_start") or (date.today() - timedelta(days=30)).isoformat()
        for key in ("date_start", "date_end", "before_date_start", "before_date_end"):
            if slots.get(key):
                date.fromisoformat(slots[key])
        if slots.get("date_start") and slots.get("date_end") and slots["date_start"] > slots["date_end"]:
            raise DomainBlocked("起始日期晚于结束日期，请修改条件", "waiting_user")
        if slots["two_dates"] and not (slots.get("before_date_start") and slots.get("before_date_end")):
            raise DomainBlocked("变化任务需要明确前期和后期两个日期范围", "waiting_user")
        return {"slots": slots}

    def resolve_aoi(self):
        from .views import _normalize_agent_bbox, resolve_district_bbox
        requested = self.context.get("input", {})
        if requested.get("bbox"):
            bbox = _normalize_agent_bbox(requested["bbox"])
            if not bbox:
                raise DomainBlocked("框选范围无效，请重新框选", "waiting_user")
            aoi = {"bbox": bbox, "polygon": requested.get("polygon"), "scope": "user_bbox"}
        else:
            place = self.context["slots"].get("place_name")
            if not place:
                raise DomainBlocked("请提供行政区名称或框选范围", "waiting_user")
            self.before_request()
            location = resolve_district_bbox(place)
            if not location.get("polygon"):
                raise DomainBlocked("行政区边界缺失，不能计算行政区比例，请提供边界或改为明确框选任务", "waiting_user")
            aoi = {"bbox": location["bbox"], "polygon": location["polygon"], "scope": "administrative_polygon", "name": location.get("name")}
        return {"aoi": aoi}

    def declare_data_requirements(self):
        slots = self.context["slots"]
        requirements = build_data_requirements(task=slots["task"], visual_detail=True,
                    spectral_bands=bool(slots["indices"]), required_indices=slots["indices"],
                    recent_acquisition=slots.get("recent_acquisition", False), two_dates=slots["two_dates"],
                    physical_measurement=slots["physical_measurement"], area_ratio=slots["area_ratio"],
                    minimum_valid_pixel_ratio=0.6, aoi_geometry_required=self.context["aoi"]["scope"] == "administrative_polygon")
        return {"requirements": requirements}

    def match_source_capabilities(self):
        source = self.context["slots"]["source"]
        requirements = self.context["requirements"]
        capabilities = SOURCE_CAPABILITIES[source]
        missing = []
        if source in ("mapbox", "tianditu", "esri"):
            missing = [key for key in ["needs_spectral_bands", "needs_two_dates", "needs_physical_measurement", "needs_area_ratio", "needs_recent_acquisition"] if requirements[key]]
        if missing:
            raise DomainBlocked("高清底图(Mapbox/天地图/Esri)缺少可追溯日期、校准波段或物理测量能力，请选择 Sentinel-2 或专用数据源", "not_supported")
        if requirements["needs_physical_measurement"]:
            raise DomainBlocked("当前输出是重采样筛查网格，原生分辨率物理测量需要专用测量数据与流程", "not_supported")
        return {"capabilities": {"status": "matched", "source": source, "required": requirements, "available": capabilities}}

    def search_scenes(self):
        slots, aoi = self.context["slots"], self.context["aoi"]
        if slots["source"] in ("mapbox", "tianditu", "esri"):
            from .agent.tools import _tool_fetch_mapbox_imagery
            ctx = {"slots": {**slots, "bbox": aoi["bbox"]}, "goal": self.run.goal, "bbox": aoi["bbox"], "mode": self.run.mode}
            self.before_request()
            result = _tool_fetch_mapbox_imagery(ctx, {"basemap_source": slots["source"]})
            if result.get("status") != "ok":
                raise DomainBlocked(result.get("message") or "高清底图影像获取失败", "external_service_unavailable")
            return {"candidates": [], "mapbox_scene_id": ctx["scene_id"], "mapbox_file_name": ctx["file_name"]}
        periods = [(slots["date_start"], slots["date_end"])]
        if slots["two_dates"]:
            periods.insert(0, (slots["before_date_start"], slots["before_date_end"]))
        collection = SOURCE_TO_COLLECTION.get(slots["source"], "sentinel-2-l2a")
        candidates = []
        for start, end in periods:
            self.before_request()
            results = get_provider_for_collection(collection).search(aoi["bbox"], start_date=start, end_date=end, max_cloud=30, limit=30, collection=collection)
            candidates.append([candidate.as_dict() for candidate in results])
        return {"candidates": candidates}

    def quality_gate(self):
        if self.context["slots"]["source"] in ("mapbox", "tianditu", "esri"):
            from .models import ImageryScene
            scene = ImageryScene.objects.get(pk=self.context["mapbox_scene_id"])
            return {"quality": {"passed": True, "source": self.context["slots"]["source"], "limitations": ["拍摄时间与传感器原生 GSD 未知，仅可视觉参考"], "scene_id": scene.id}}
        selected = []
        bbox = self.context["aoi"]["bbox"]
        xs = np.linspace(bbox["min_lng"], bbox["max_lng"], 66)[1:-1]
        ys = np.linspace(bbox["min_lat"], bbox["max_lat"], 66)[1:-1]
        x, y = np.meshgrid(xs, ys)
        slots = self.context["slots"]
        collection = SOURCE_TO_COLLECTION.get(slots["source"], "sentinel-2-l2a")
        profile = get_collection_profile(collection)
        # 必需资产按 profile 生成:S2 保持原波段+SCL;landsat 用 qa_pixel 位掩膜;
        # SAR 只要求 vv;DEM 只要求 data。
        band_aliases = profile.get("band_aliases") or {}
        def _band_asset(band):
            return band_aliases.get(band, "swir16" if band == "swir" else band)
        if profile.get("cloud_property") and profile.get("qa_kind") == "bitmask":
            required = {"red", "green", "blue", "nir08", "qa_pixel"}
            for index in self.context["requirements"]["required_indices"]:
                required.update(_band_asset(band) for band in INDEX_DEFINITIONS[index]["bands"])
        elif profile.get("cloud_property"):
            required = {"visual", "scl", "red", "green", "blue"}
            for index in self.context["requirements"]["required_indices"]:
                required.update("swir16" if band == "swir" else band for band in INDEX_DEFINITIONS[index]["bands"])
        else:
            required = {profile["render"]["asset"]}
        cloud_exempt = profile.get("cloud_property") is None
        for position, candidates in enumerate(self.context["candidates"]):
            groups = defaultdict(list)
            start = slots["before_date_start"] if slots["two_dates"] and position == 0 else slots["date_start"]
            end = slots["before_date_end"] if slots["two_dates"] and position == 0 else slots["date_end"]
            for candidate in candidates:
                acquired = candidate.get("acquired_at")
                cloud = candidate.get("cloud_percent")
                cloud_ok = True if cloud_exempt else (cloud is not None and 0 <= float(cloud) <= 30)
                if (not acquired or not start <= acquired[:10] <= end or not cloud_ok
                        or not required.issubset(candidate.get("assets", {}))):
                    continue
                groups[acquired[:10]].append(candidate)
            found = None
            for day in sorted(groups, reverse=True):
                group = sorted(groups[day], key=lambda candidate: 0.0 if candidate.get("cloud_percent") is None else candidate["cloud_percent"])
                coverage = np.zeros(x.shape, dtype=bool)
                for candidate in group:
                    bound = candidate["bbox"]
                    coverage |= (x >= bound["min_lng"]) & (x <= bound["max_lng"]) & (y >= bound["min_lat"]) & (y <= bound["max_lat"])
                if coverage.mean() >= 0.98:
                    found = group
                    break
            if not found:
                raise DomainBlocked("没有同时满足日期、云量、必需资产和覆盖率的单景或同日场景，请修改条件", "waiting_user")
            selected.append(found)
        if len(selected) == 2 and selected[0][0]["acquired_at"][:10] >= selected[1][0]["acquired_at"][:10]:
            raise DomainBlocked("前后两期必须是严格递增的独立日期，不能复用同日或跨日期拼接结果")
        return {"quality": {"passed": True, "max_cloud_percent": 30, "minimum_footprint_coverage": 0.98,
                            "coverage_method": "64x64 bbox footprint sample; pixel coverage verified after asset read"}, "selected_periods": selected}

    def select_product(self):
        periods = self.context.get("selected_periods", [])
        return {"periods": periods, "selection_policy": "latest_qualified_date_per_period; same_date_only"}

    def read_assets(self):
        if self.context["slots"]["source"] in ("mapbox", "tianditu", "esri"):
            from .media_paths import safe_media_path
            from django.conf import settings
            from pathlib import Path
            path = safe_media_path(str(Path(settings.MEDIA_ROOT) / "satellite_imgs"), self.context["mapbox_file_name"], (".jpg", ".png", ".jpeg"))
            if not path:
                raise DomainBlocked("高清底图影像文件缺失")
            from PIL import Image
            with Image.open(path) as image:
                image.verify()
            product = write_product(self.claim, "jpg", Path(path).read_bytes())
            self.save_file(product, "高清底图视觉参考", "imagery", "image/jpeg")
            return {"assets": product, "visual_file": product, "grid_contract": {"source": self.context["slots"]["source"], "temporal": "unknown"}}
        collection = SOURCE_TO_COLLECTION.get(self.context["slots"]["source"], "sentinel-2-l2a")
        arrays, contract = read_scene_assets(self.context["periods"], self.context["aoi"]["bbox"], self.context["requirements"]["required_indices"], before_request=self.before_request, collection=collection)
        product = save_arrays(self.claim, arrays)
        self.save_file(product, "校准波段", "raster_bundle", "application/octet-stream")
        return {"assets": product, "grid_contract": contract}

    def apply_aoi_and_qa_mask(self):
        if self.context["slots"]["source"] in ("mapbox", "tianditu", "esri"):
            return {"mask_statistics": [], "mask_applicability": "visual_only_no_physical_measurement"}
        arrays, stats = apply_masks(load_arrays(self.run.id, self.context["assets"]), self.context["grid_contract"], self.context["aoi"]["polygon"])
        product = save_arrays(self.claim, arrays)
        self.save_file(product, "AOI 与 QA 掩膜后波段", "raster_bundle", "application/octet-stream")
        return {"masked_assets": product, "mask_statistics": stats}

    def compute_metric(self):
        if self.context["slots"]["source"] in ("mapbox", "tianditu", "esri"):
            return {"metrics": [], "metric_applicability": "visual_only_no_spectral_measurement"}
        arrays, summaries = calculate_products(load_arrays(self.run.id, self.context["masked_assets"]), self.context["requirements"]["required_indices"], len(self.context["periods"]))
        product = save_arrays(self.claim, arrays)
        self.save_file(product, "指数计算数据", "raster_bundle", "application/octet-stream")
        for summary in summaries:
            period, index = summary["period"], summary["index"]
            raster = write_product(self.claim, f"p{period}-{index}.tif", geotiff_bytes(arrays[f"p{period}_{index}"], self.context["aoi"]["bbox"]))
            asset_id = self.save_file(raster, f"时相 {period + 1} · {index.upper()}", "index_raster", "image/tiff")
            scene_id = "+".join(candidate["item_id"] for candidate in self.context["periods"][period])
            if len(scene_id) > 130:
                from .run_journal import digest
                scene_id = "mosaic:" + digest(scene_id)
            item = _evidence(self.claim, kind="computed_metric", scene_id=scene_id, metric=index, value=summary,
                             method=summary["method"], contract=self.context["grid_contract"],
                             mask=self.context["mask_statistics"][period], aoi=self.context["aoi"], asset_id=asset_id, limitations=summary["limitations"])
            self.evidence.append(item)
        return {"metric_arrays": product, "metrics": summaries}

    def compare_temporal_scenes(self):
        if not self.context["requirements"]["needs_two_dates"]:
            return None
        if not self.context["requirements"]["required_indices"]:
            raise DomainBlocked("两期变化需要明确可计算的指数", "not_supported")
        arrays, summaries = compare_products(load_arrays(self.run.id, self.context["metric_arrays"]),
                    self.context["requirements"]["required_indices"], self.context["mask_statistics"][0]["aoi_pixel_count"])
        for index, values in arrays.items():
            product = write_product(self.claim, f"{index}-change.tif", geotiff_bytes(values, self.context["aoi"]["bbox"]))
            asset_id = self.save_file(product, f"{index.upper()} 两期变化", "change_raster", "image/tiff")
            item = _evidence(self.claim, kind="temporal_change", scene_id="before+after", metric=index,
                             value=next(summary for summary in summaries if summary["index"] == index),
                             method="after-before on common QA grid", contract=self.context["grid_contract"],
                             mask={"common_mask": True}, aoi=self.context["aoi"], asset_id=asset_id)
            self.evidence.append(item)
        return {"changes": summaries}

    def _current_evidence(self):
        return list(self.run.evidence_v2.filter(evidence_id__startswith=f"v{self.run.plan_version}:").values())

    def visual_review(self):
        if self.context["slots"]["source"] in ("mapbox", "tianditu", "esri"):
            product = self.context["visual_file"]
            image = read_product(self.run.id, product)
        else:
            image = rgb_preview(load_arrays(self.run.id, self.context["masked_assets"]), len(self.context["periods"]) - 1)
            product = write_product(self.claim, "png", image)
            self.save_file(product, "当前场景真彩色复核图", "imagery", "image/png")
        # Data facts are available even in visual-only tasks; visual inference is
        # saved separately, never presented as an independently measured metric.
        fact = _evidence(self.claim, kind="data_fact", scene_id=str(self.context.get("mapbox_scene_id") or "selected-scenes"),
                         value={"periods": self.context["periods"], "quality": self.context["quality"]},
                         contract=self.context["grid_contract"], aoi=self.context["aoi"], method="source_metadata")
        refs = [item["evidence_id"] for item in self._current_evidence()] + [fact["evidence_id"]]
        decision = self.provider().review({"goal": self.run.goal, "step": "visual_review", "metrics": self.context["metrics"],
                "data_contract": {key: value for key, value in self.context["grid_contract"].items() if key != "scenes"},
                "evidence_refs": refs, "instruction": "核对真彩色图与计算结果，只描述可见事实和明确限制。返回 final 并引用给定证据。"},
                images=["data:image/png;base64," + base64.b64encode(image).decode("ascii")])
        if decision["type"] != "final" or not set(decision["evidence_refs"]).issubset(refs):
            raise DomainBlocked("视觉复核未形成有效证据，需重试或修改条件", "waiting_user")
        self.evidence.extend([fact, _evidence(self.claim, kind="model_inference", scene_id="selected-scenes", value=decision["content"],
                                             method=f"{self.run.provider}/{self.run.model}", aoi=self.context["aoi"],
                                             contract=self.context["grid_contract"], limitations=decision["limitations"])])
        return {"visual_review": decision, "visual_file": product}

    def evidence_compilation(self):
        items = self._current_evidence()
        required = self.context["requirements"]["required_indices"]
        for index in required:
            if sum(item["metric"] == index and item["kind"] == "computed_metric" for item in items) != len(self.context["periods"]):
                raise DomainBlocked("必需指标证据不完整，不能生成结论")
        if not any(item["kind"] == "model_inference" for item in items):
            raise DomainBlocked("缺少视觉复核证据")
        if self.context["requirements"]["needs_two_dates"] and not any(item["kind"] == "temporal_change" for item in items):
            raise DomainBlocked("缺少两期变化证据")
        return {"evidence_refs": [item["evidence_id"] for item in items]}

    def final_review(self):
        evidence = [{key: item[key] for key in ("evidence_id", "kind", "metric", "value", "limitations")} for item in self._current_evidence()]
        decision = self.provider().review({"goal": self.run.goal, "requirements": self.context["requirements"],
                                         "evidence": evidence, "evidence_refs": self.context["evidence_refs"],
                                         "instruction": "输出正式复核结论。必须引用全部必需指标、变化和视觉证据，并明确限制。"})
        if decision["type"] != "final" or set(decision["evidence_refs"]) != set(self.context["evidence_refs"]):
            raise DomainBlocked("最终回答未完整引用本次任务证据，不能标记为完成", "waiting_user")
        return {"final": decision}

    def publish_artifacts(self):
        from docx import Document
        final = self.context["final"]
        document = Document()
        document.add_heading("遥感调查报告", 0)
        document.add_paragraph(self.run.goal)
        document.add_heading("复核结论", 1)
        document.add_paragraph(final["content"])
        document.add_heading("限制", 1)
        for limitation in final["limitations"]:
            document.add_paragraph(limitation, style="List Bullet")
        document.add_heading("证据索引", 1)
        for item in self._current_evidence():
            document.add_heading(item["evidence_id"], 2)
            document.add_paragraph(f"类型：{item['kind']}；方法：{item['method']}")
            document.add_paragraph(json.dumps({key: item[key] for key in ("scene_id", "asset_id", "metric", "value", "aoi", "mask_statistics", "data_contract", "limitations")}, ensure_ascii=False))
        buf = BytesIO()
        document.save(buf)
        product = write_product(self.claim, "docx", buf.getvalue())
        self.save_file(product, "遥感调查报告", "report", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", refs=final["evidence_refs"])
        result = {"goal": self.run.goal, "final": final, "requirements": self.context["requirements"],
                  "metrics": self.context["metrics"], "changes": self.context.get("changes", []), "evidence_refs": final["evidence_refs"]}
        json_product = write_product(self.claim, "json", json.dumps(result, ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8"))
        self.save_file(json_product, "结构化调查结果", "final_result", "application/json", refs=final["evidence_refs"])
        return {"report": product, "result": result}


@transaction.atomic
def project_legacy(run_id):
    """Compatibility is one-way for DAG runs: database run -> legacy session."""
    run = AgentRun.objects.get(pk=run_id)
    if run.execution_engine != "dag":
        return
    session = AgentSession.objects.select_for_update().filter(artifacts__run_id=run_id).first()
    if not session:
        return
    context = load_checkpoint(run_id)["context"]
    session.status = "running" if run.status in {"queued", "planning", "running", "retrying"} else ("failed" if run.status in {"cancelled", "cancelling", "blocked", "not_supported", "external_service_unavailable"} else run.status)
    session.goal, session.mode, session.error = run.goal, run.mode, run.error
    session.cancel_requested = run.status in {"cancelling", "cancelled"}
    session.slots = context.get("slots", {})
    artifacts = {**(session.artifacts or {}), "run_id": run.id, "execution_mode": "queue"}
    artifacts.pop("worker_claim", None)
    artifacts.pop("worker_claimed_at", None)
    if run.status == "waiting_user" or run.error:
        artifacts["waiting"] = {"message": run.error or "已暂停，可继续执行或修改条件", "options": [{"code": "retry_step", "label": "重试当前步骤"}, {"code": "cancel", "label": "取消任务"}], "data": {"step_id": run.current_step_id, "paused_by_user": bool(context.get("paused_by_user"))}}
    else:
        artifacts.pop("waiting", None)
    if run.status == "completed":
        artifacts.update({"final_answer": context["final"]["content"], "final_evidence_refs": context["final"]["evidence_refs"], "final_limitations": context["final"]["limitations"]})
    session.artifacts = artifacts
    session.save(update_fields=["status", "goal", "mode", "error", "cancel_requested", "slots", "artifacts", "updated_at"])


def execute_run(run_id, worker_id, *, provider_factory=configured_provider, max_steps=30):
    for _ in range(max_steps):
        claim = claim_step(run_id, worker_id, lease_seconds=600)
        if claim is None:
            run = AgentRun.objects.get(pk=run_id)
            if run.status in {"running", "retrying"} and run.steps.filter(step_id="publish_artifacts", status="completed").exists():
                complete_run(run_id)
                project_legacy(run_id)
            return
        run = AgentRun.objects.get(pk=run_id)
        context = load_checkpoint(run_id)["context"]
        handler = SatelliteHandlers(run, claim, context, provider_factory=provider_factory)
        try:
            output = getattr(handler, claim.kind)()
            if output is None:
                skip_step(claim, "当前任务没有请求两期变化")
            else:
                finish_step(claim, output, context_patch=output, evidence=handler.evidence, artifacts=handler.artifacts)
            if claim.step_id == "publish_artifacts":
                complete_run(run_id)
                return
        except StepConflict as exc:
            # Cancellation or takeover has authority over this worker.
            if exc.details["code"] != "lease_lost":
                fail_step(claim, str(exc), status="blocked")
            return
        except Exception as exc:
            if isinstance(exc, DomainBlocked):
                status, message, retryable = exc.status, str(exc), False
            elif isinstance(exc, ProviderError):
                status, message, retryable = "external_service_unavailable", exc.message, exc.retryable
            elif isinstance(exc, (requests.RequestException, TimeoutError)):
                status, message, retryable = "external_service_unavailable", "外部数据服务请求失败，请重试", True
            elif isinstance(exc, ValueError):
                status, message, retryable = "blocked", str(exc)[:500], False
            else:
                status, message, retryable = "failed", "步骤执行失败：" + type(exc).__name__, False
            try:
                fail_step(claim, message, status=status, retryable=retryable)
            except StepConflict:
                return
            # Persist a user-facing recovery contract in the legacy projection;
            # the AgentRun status remains the authoritative failure classification.
            project_legacy(run_id)
            return
        finally:
            project_legacy(run_id)


@transaction.atomic
def complete_run(run_id):
    run = lock_run(run_id)
    if run.status in {"completed", "cancelled", "cancelling"}:
        return
    from .run_acceptance import assess_run
    assessment = assess_run(run, evidence=list(run.evidence_v2.all()), artifacts=list(run.artifacts_v2.all()))
    if not assessment["passed"]:
        reason = "最终验收条件未满足：" + ", ".join(assessment["missing_requirements"])
        _set_state(run, "blocked", reason)
        append_locked(run, "run.blocked", {**assessment, "user_action": "retry_step_or_replan", "reason": reason}, command_key=uuid.uuid4().hex)
        project_legacy(run_id)
        return
    context = load_checkpoint(run_id)["context"]
    analysis = AnalysisRun.objects.create(query=run.goal, status="completed", parameters=context["requirements"],
                    scene_ids=[candidate["product_id"] for period in context["periods"] for candidate in period],
                    result=context["result"], model_versions={"provider": run.provider, "model": run.model}, finished_at=timezone.now())
    finding = Finding.objects.create(analysis_run=analysis, finding_type="remote_sensing_review", label="遥感复核结论",
                    bbox=context["aoi"]["bbox"], geometry=context["aoi"]["polygon"], source_method="calibrated_spectral_and_visual",
                    metadata={"agent_run_id": run.id, "evidence_refs": context["final"]["evidence_refs"]})
    for item in run.evidence_v2.filter(evidence_id__in=context["final"]["evidence_refs"]):
        Evidence.objects.create(finding=finding, evidence_type=item.kind, bbox=context["aoi"]["bbox"], metric=item.metric,
                                description=json.dumps({"evidence_id": item.evidence_id, "value": item.value}, ensure_ascii=False))
    context["analysis_run_id"] = analysis.id
    _set_state(run, "completed", "")
    append_locked(run, "run.completed", assessment, context=context, command_key=f"complete:{run.plan_version}")
